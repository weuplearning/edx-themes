# -*- coding: utf-8 -*-
#!/usr/bin/env python
import importlib
import sys
importlib.reload(sys)
import os
from io import BytesIO

os.environ.setdefault("DJANGO_SETTINGS_MODULE", "lms.envs.production")
os.environ.setdefault("LMS_CFG", "/edx/etc/lms.yml")
os.environ.setdefault("lms.envs.production,SERVICE_VARIANT", "lms")
os.environ.setdefault(
    "PATH",
    "/edx/app/edxapp/venvs/edxapp/bin:/edx/app/edxapp/edx-platform/bin:"
    "/edx/app/edxapp/.rbenv/bin:/edx/app/edxapp/.rbenv/shims:"
    "/edx/app/edxapp/.gem/bin:/edx/app/edxapp/edx-platform/node_modules/.bin:"
    "/usr/local/sbin:/usr/local/bin:/usr/sbin:/usr/bin:/sbin:/bin"
)
os.environ.setdefault("SERVICE_VARIANT", "lms")
os.chdir("/edx/app/edxapp/edx-platform")

from django.core.wsgi import get_wsgi_application
application = get_wsgi_application()


#############################################################
#         ^ SETUP ENVIRONNEMENT VARIABLE FOR KOA ^          #
#                START BEYOND THIS LINE                     #
#############################################################

import json
import time
import smtplib
import logging

from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email.mime.base import MIMEBase
from email import encoders

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from opaque_keys.edx.locator import CourseLocator
from common.djangoapps.student.models import CourseEnrollment
from lms.djangoapps.courseware.courses import get_course_by_id
from lms.djangoapps.wul_apps.best_grade.helpers import check_best_grade
from lms.djangoapps.wul_apps.models import WulCourseEnrollment


log = logging.getLogger()


emails = sys.argv[1].split(";")
course_ids = sys.argv[2].split(";")


all_users_data = {}

headers = [
    'Email', 'Nom complet', 'Adresse', 'Code postal', 'Ville', 'Région',
    'Pays si pas FR', 'Profession', 'Profession si autre',
    'E1 - Temps passé', 'E1 - Progression',
    'E2 - Temps passé', 'E2 - Progression',
    'GDA - Temps passé', 'GDA - Progression',
    'E1 Bucco - Temps passé', 'E1 Bucco - Progression',
    'E1 EN - Temps passé', 'E1 EN - Progression',
]

course_column_map = {
    "E1": "course-v1:icope+e1+2026",
    "E2": "course-v1:icope+ea+2026",
    "GDA": "course-v1:icope+gda+2026",
    "E1 Bucco": "course-v1:icope+soins_bucco_dentaires+2024",
    "E1 EN": "course-v1:icope_en+1+2026",
}




for course_id in course_ids:
    log.info('------------> Start calculate grades for course : ' + str(course_id))

    course_key = CourseLocator.from_string(course_id)
    course = get_course_by_id(course_key)
    course_enrollments = CourseEnrollment.objects.filter(course_id=course_key)

    for enrollment in course_enrollments:
        user = enrollment.user

        if str(user.email).find('@yopmail') != -1 or str(user.email).find('@weuplearning') != -1 or str(user.email).find('@themoocagency') != -1 :
            continue

        # TimeTracking
        try:
            wul_course_enrollment = WulCourseEnrollment.objects.get(
                course_enrollment_edx__user=user,
                course_enrollment_edx__course_id=course_key,
            )
            time_total_course = wul_course_enrollment.global_time_tracking
        except Exception:
            time_total_course = 0

        if time_total_course == 0:
            continue

        user_id = str(user.id)

        if user_id not in all_users_data:
            not_found_str = 'n.a.'

            try:
                custom_field = json.loads(user.profile.custom_field)
            except Exception:
                custom_field = {}

            user_data = {
                "email": user.email,
                "name": user.profile.name,
                "adress": custom_field.get('adress', not_found_str),
                "post_code": custom_field.get('post_code', not_found_str),
                "city": custom_field.get('city', not_found_str),
                "region": custom_field.get('region', not_found_str),
                "country_icopeen": custom_field.get('country_icopeen', not_found_str),
                "profession": custom_field.get('profession', not_found_str),
                "profession_autre": custom_field.get('profession_autre', not_found_str),
            }

            all_users_data[user_id] = {
                "profil": user_data,
                "courses": {},
            }

        # Progression / meilleure note pour ce cours
        try:
            grades_test = check_best_grade(user, course, force_best_grade=True)
            progression = round(grades_test.summary['percent'] * 100, 2)
        except Exception:
            progression = 0

        all_users_data[user_id]["courses"][course_id] = {
            "time_total": time_total_course,
            "progression": progression,
        }


# WRITE EXCEL
timestr = time.strftime("%Y_%m_%d")
wb = Workbook()
sheet = wb.active
sheet.title = 'Rapport de notes'
filename = '/home/edxtma/csv/Icope_en_fr_{}.xlsx'.format(timestr)

# En-têtes sur la première ligne
for col_index, header in enumerate(headers, start=1):
    cell = sheet.cell(row=1, column=col_index, value=header)
    cell.fill = PatternFill("solid", fgColor="6B9AAF")
    cell.font = Font(b=False, color="FFFFFF")

# Une ligne par utilisateur, à partir de la ligne 2
row_index = 2

for user_id, user_entry in all_users_data.items():
    profil = user_entry['profil']
    courses = user_entry['courses']

    profil_values = [
        profil.get('email', 'n.a.'),
        profil.get('name', 'n.a.'),
        profil.get('adress', 'n.a.'),
        profil.get('post_code', 'n.a.'),
        profil.get('city', 'n.a.'),
        profil.get('region', 'n.a.'),
        profil.get('country_icopeen', 'n.a.'),
        profil.get('profession', 'n.a.'),
        profil.get('profession_autre', 'n.a.'),
    ]

    # Colonnes 1 à 9 : informations utilisateur
    for col_index, value in enumerate(profil_values, start=1):
        sheet.cell(row=row_index, column=col_index, value=value)

    # Colonnes 10 à 19 : données par cours
    col_index = len(profil_values) + 1

    for course_label, mapped_course_id in course_column_map.items():
        course_data = courses.get(mapped_course_id, {})

        # Temps passé en minutes
        time_total = course_data.get('time_total', 'n.a.')
        if isinstance(time_total, (int, float)):
            time_total_min = f"{time_total} min"
        else:
            time_total_min = time_total   
        sheet.cell(row=row_index, column=col_index, value=time_total_min)

        # Progression
        progression = course_data.get('progression', 'n.a.')
        if isinstance(progression, (int, float)):
            progression_value = f"{progression}%"
        else:
            progression_value = progression
        sheet.cell(row=row_index, column=col_index + 1, value=progression_value)

        col_index += 2

    row_index += 1


# Ajustement automatique simple de la largeur des colonnes
for column_cells in sheet.columns:
    max_length = 0

    for cell in column_cells:
        if cell.value is not None:
            max_length = max(max_length, len(str(cell.value)))

    sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max_length + 2


output = BytesIO()
wb.save(output)
_files_values = output.getvalue()


# SEND MAILS
html = (
    "<html><head></head><body>"
    "<p>Bonjour,<br/><br/>"
    "Vous trouverez en pièce jointe le rapport de temps passé pour les différents cours "
    "icope-formation, FR et EN.<br/><br/>"
    "Bonne r&eacute;ception<br/>"
    "L'&eacute;quipe WeUp Learning"
    "</p></body></html>"
)


fromaddr = "ne-pas-repondre@themoocagency.com"

for email in emails:
    msg = MIMEMultipart()
    msg['From'] = fromaddr
    msg['To'] = email
    msg['Subject'] = "Rapport de temps passé - MOOC icope"

    # Pièce jointe Excel
    attachment = _files_values
    part = MIMEBase('application', 'octet-stream')
    part.set_payload(attachment)
    encoders.encode_base64(part)
    part.add_header(
        'Content-Disposition',
        'attachment; filename="%s"' % os.path.basename(filename)
    )
    msg.attach(part)

    # Corps HTML
    part2 = MIMEText(html.encode('utf-8'), 'html', 'utf-8')
    msg.attach(part2)

    server = smtplib.SMTP('mail3.themoocagency.com', 25)
    server.starttls()
    server.login('contact', 'waSwv6Eqer89')
    server.sendmail(fromaddr, email, msg.as_string())
    server.quit()

    log.info('Email sent to ' + str(email))


log.info('------------> Finish calculate grades and write xlsx report')

# sudo /edx/app/edxapp/venvs/edxapp/bin/python /edx/app/edxapp/edx-themes/icope/lms/utils/grade_report_en_fr.py "cyril.adolf@weuplearning.com" "course-v1:icope+e1+2026;course-v1:icope+gda+2026;course-v1:icope+ea+2026;course-v1:icope+soins_bucco_dentaires+2024;course-v1:icope_en+1+2026"

